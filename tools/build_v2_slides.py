"""Teacher-facing decks for lessons 2–6, using the school's Lesson 1 theme.

Slides are written to the course workspace first. Copy the five final files to
the matching Grade 8 lesson directories after review.
"""
from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT.parent / 'CPP8_Bai01_MachMoi' / 'output' / 'Slide_Bai01_RaLenh_TheoMachDay_20261003_FINAL.pptx'
SCRATCH = Path(r'D:\OneDrive - Lawrence S.Ting School\General\2026-2027\02. Tài liệu giảng dạy\Grade 8\03. Lap trinh C++\_Tools\_Web\assets\scratch')
OUT = ROOT.parent / 'CPP8_NewSlides'
OUT.mkdir(exist_ok=True)
LESSONS = json.loads((ROOT / 'tools' / 'v2-lessons.json').read_text(encoding='utf-8'))

INK = RGBColor(22, 51, 72)
NAVY = RGBColor(9, 29, 51)
TEAL = RGBColor(0, 127, 143)
PALE = RGBColor(236, 247, 249)
GREEN = RGBColor(230, 247, 238)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(81, 103, 118)
ORANGE = RGBColor(255, 234, 195)

HOOK = {
    2: ('Cửa hàng bán vé', 'Mỗi vé 12 000 đồng. Ba vé hết bao nhiêu?', 'cout << 12000 + 12000 + 12000;', '36 000', 'Nếu đổi thành 4 vé, con phải sửa mấy chỗ?'),
    3: ('Robot nói gì cũng như nhau', 'Dù điểm là 4 hay 8, robot đều in cùng một câu.', 'cout << "Con yeu thay Vuong!";', 'Con yeu thay Vuong!', 'Robot cần biết điều gì trước khi chọn câu trả lời?'),
    4: ('Một việc làm 100 lần', 'Thầy muốn máy in “Tôi thích C++” 100 lần.', 'cout << "Toi thich C++" << endl;\ncout << "Toi thich C++" << endl;\ncout << "Toi thich C++" << endl;', 'Ba dòng giống nhau', 'Nếu làm 100 lần, con sẽ viết thế nào?'),
    5: ('Một lời chào ở ba nơi', 'Robot chào lúc mở màn, sau tiết mục và lúc kết thúc.', 'cout << "XIN CHAO" << endl;\n// ...\ncout << "XIN CHAO" << endl;', 'Lời chào lặp lại', 'Nếu muốn đổi lời chào, con phải tìm và sửa mấy chỗ?'),
    6: ('Máy chạy nhưng kết quả sai', 'Bạn An được 5 điểm. Máy lại in THU LAI.', 'if (diem > 5) cout << "DAT";\nelse cout << "THU LAI";', 'THU LAI', 'Máy sai ở cú pháp hay ở điều kiện?'),
}

SOLUTION = {
    2: '#include <iostream>\nusing namespace std;\nint main() {\n  int gia, so;\n  cin >> gia >> so;\n  cout << "Tong tien: " << gia * so;\n  return 0;\n}',
    3: '#include <iostream>\nusing namespace std;\nint main() {\n  int h; cin >> h;\n  if (h >= 140) cout << "DUOC VAO";\n  else cout << "CHUA DU CAO";\n  return 0;\n}',
    4: '#include <iostream>\nusing namespace std;\nint main() {\n  int n; cin >> n;\n  for (int i = 1; i <= n; i++) {\n    cout << "CO GANG " << i << endl;\n  }\n  return 0;\n}',
    5: '#include <iostream>\nusing namespace std;\nvoid chao() { cout << "XIN CHAO" << endl; }\nvoid vongTay() { cout << "* * *" << endl; }\nint main() {\n  chao(); vongTay(); chao();\n  return 0;\n}',
    6: '#include <iostream>\nusing namespace std;\nint main() {\n  int keo, ban; cin >> keo >> ban;\n  cout << keo / ban << endl;\n  cout << keo % ban;\n  return 0;\n}',
}

def shape(slide, x, y, w, h, fill, line=None, rounded=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line: s.line.color.rgb = line; s.line.width = Pt(1)
    else: s.line.fill.background()
    if rounded:
        try: s.adjustments[0] = .1
        except Exception: pass
    return s

def txt(slide, value, x, y, w, h, size=22, color=INK, bold=False, font='Arial', align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for n, line in enumerate(str(value).split('\n')):
        p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        p.text = line; p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
        p.space_after = Pt(9); p.line_spacing = 1.06
        if align is not None: p.alignment = align
    return box

def make_deck(n, data):
    prs = Presentation(TEMPLATE)
    blue_logo = next(sh.image.blob for sh in prs.slides[0].shapes if sh.shape_type == 13 and sh.left > Inches(10))
    white_logo = None
    for old in prs.slides:
        for sh in old.shapes:
            if sh.shape_type == 13 and sh.left > Inches(10) and sh.image.blob != blue_logo:
                white_logo = sh.image.blob; break
        if white_logo: break
    for sid in list(prs.slides._sldIdLst):
        prs.part.drop_rel(sid.rId); prs.slides._sldIdLst.remove(sid)

    def slide(title, section, dark=False):
        s = prs.slides.add_slide(prs.slide_layouts[6]); shape(s, 0, 0, 13.333, 7.5, NAVY if dark else WHITE); shape(s, 0, 0, 13.333, .055, TEAL)
        txt(s, f'C++ LỚP 8 · BÀI {n}  /  {section}', .85, .26, 10.5, .3, 10, RGBColor(137,224,223) if dark else TEAL, True)
        txt(s, title, .79, .77, 11.65, .77, 30, WHITE if dark else INK, True)
        logo = white_logo if dark and white_logo else blue_logo
        s.shapes.add_picture(BytesIO(logo), Inches(11.9), Inches(.2), width=Inches(.79))
        return s

    def card(s, title, body, x, y, w, h, fill=PALE, fontsize=20):
        shape(s, x, y, w, h, fill, RGBColor(205, 225, 231), True)
        txt(s, title, x+.18, y+.16, w-.36, .42, 16, TEAL, True)
        txt(s, body, x+.2, y+.64, w-.4, h-.78, fontsize, INK)

    def code(s, body, x, y, w, h, size=18):
        shape(s, x, y, w, h, NAVY, None, True); shape(s, x, y, w, .43, RGBColor(31, 67, 87))
        txt(s, '●  ●  ●      main.cpp', x+.15, y+.08, w-.3, .24, 11, RGBColor(210,235,242), True, 'Consolas')
        txt(s, body, x+.22, y+.59, w-.44, h-.7, size, WHITE, False, 'Consolas')

    hook_title, scenario, hook_code, hook_result, hook_question = HOOK[n]
    s=slide(data['title'].upper(),'MỞ BÀI · 4 PHÚT',True)
    txt(s,scenario,.93,2.05,7.7,1.35,28,WHITE,True)
    txt(s,hook_question,.93,4.18,8.0,1.3,23,RGBColor(160,232,223))
    shape(s,9.5,2.07,2.2,2.2,RGBColor(26,94,108),None,True);txt(s,data['emoji'],9.98,2.41,1.25,1.2,63,WHITE,True,align=PP_ALIGN.CENTER)

    s=slide('Đoán trước khi chạy','HOOK · CÂU HỎI CHO CẢ LỚP',True)
    code(s,hook_code,.86,1.75,7.3,3.55,19)
    txt(s,'1. Chương trình sẽ in gì?\n2. Nếu thay dữ liệu, con sửa ở đâu?',8.45,2.05,3.9,2.2,22,WHITE,True)
    txt(s,'Cho học sinh đoán cá nhân, rồi giải thích với bạn bên cạnh.',.9,5.85,11.6,.5,18,RGBColor(150,225,224))

    s=slide('Chạy thử và đặt vấn đề','DẪN VÀO BÀI')
    card(s,'Kết quả hiện ra',hook_result,.85,1.79,5.7,3.45,GREEN,25)
    card(s,'Điều cần giải quyết',hook_question,6.8,1.79,5.7,3.45,PALE,22)
    txt(s,'Hôm nay: '+data['scope'],.88,5.63,11.5,.95,22,TEAL,True)

    s=slide('Từ Scratch sang C++' if n<6 else 'Từ chạy thử đến kiểm thử','KẾT NỐI KIẾN THỨC CŨ · 3 PHÚT')
    image = SCRATCH / (f'bai{n:02d}_slide.png' if n in (2,3,5) else f'bai{n:02d}.png')
    if image.exists():
        scratch_width = 3.5 if n == 5 else 4.55
        s.shapes.add_picture(str(image), Inches(.85), Inches(1.82), width=Inches(scratch_width))
        if n==2: bridge='Khối + và “chia lấy dư” trong Scratch tương ứng với + và % trong C++.\n\nBiến và cin sẽ xuất hiện ở hai trạm đầu trên web.'
        elif n==3: bridge='Khối nếu / nếu không có hai nhánh giống if / else.\n\nHình dùng > 5. Với >= 5, điểm đúng 5 sẽ đi nhánh nào?'
        else: bridge=data['scratch']
        card(s,'Cách nghĩ chung',bridge,6.0,1.8,6.42,3.75,GREEN,19)
        txt(s,'Hỏi: Trong hai cách viết, thao tác nào giống nhau?',.88,6.23,11.6,.5,20,TEAL,True)
    else:
        card(s,'Dự đoán','Chọn đầu vào.\nNói kết quả con chờ đợi.',.9,1.8,3.6,3.72,GREEN,21)
        card(s,'Chạy thử','Chạy chương trình.\nĐọc kết quả thật và thông báo lỗi.',4.84,1.8,3.6,3.72,PALE,21)
        card(s,'Sửa lỗi','So sánh hai kết quả.\nSửa một chỗ rồi chạy lại.',8.78,1.8,3.6,3.72,GREEN,21)
        txt(s,'Hỏi: Chạy được có nghĩa là đã đúng yêu cầu chưa?',.88,6.1,11.6,.5,20,TEAL,True)

    s=slide('Bốn việc con sẽ làm','MỤC TIÊU BÀI HỌC')
    for i,stage in enumerate(data['stages']):
        x=.86+(i%2)*6.27;y=1.75+(i//2)*2.22
        card(s,f'{i+1}. {stage["name"]}',stage['lesson'],x,y,5.96,2.02,PALE if i%2 else GREEN,18)

    s=slide(f'Mở web: {data["title"]}','CÁ NHÂN · 17 PHÚT')
    card(s,'Đọc','Mỗi trạm có một ý mới, ví dụ C++ và câu hỏi ngắn.',.79,1.72,3.82,3.7)
    card(s,'Chơi',data['mechanic']+'\nSau đó tự viết và chạy C++.',4.74,1.72,3.82,3.7,GREEN)
    card(s,'Dừng','Qua trạm cuối, tải phiếu PDF.\nDừng web, chờ hỏi đáp cả lớp.',8.68,1.72,3.82,3.7)
    txt(s,f'Đường dẫn: vuongndlst.github.io/cpp-8/bai{n:02d}/game.html',.86,5.96,11.6,.53,19,TEAL,True)

    s=slide('Dừng web: em hiểu gì?','CLASSPOINT · HỎI ĐÁP · 4 PHÚT',True)
    txt(s,'1. '+data['discuss'][0],.95,2.08,11.4,1.1,25,WHITE,True)
    txt(s,'2. '+data['discuss'][1],.95,3.72,11.4,1.18,25,WHITE,True)
    txt(s,'Mời học sinh chỉ một dòng mã hoặc kết quả đã chạy để giải thích.',.96,5.74,11.2,.53,18,RGBColor(153,226,224))

    s=slide('Ghi bài: ý chính thứ nhất','CHỐT KIẾN THỨC · 3 PHÚT')
    card(s,'Ghi vào vở','\n'.join(data['note'][:2]),.85,1.75,6.0,4.35,GREEN,21)
    card(s,'Ý nghĩa',data['concept'],7.08,1.75,5.45,4.35,PALE,19)

    s=slide('Ghi bài: cú pháp và ví dụ','CHỐT KIẾN THỨC · 3 PHÚT')
    code(s,data['stages'][1]['example'],.86,1.74,6.52,3.4,23)
    card(s,'Nhớ khi viết','\n'.join(data['note'][2:]),7.64,1.75,4.82,3.98,GREEN,19)
    txt(s,'Đọc mã từ trên xuống dưới. Dự đoán rồi mới chạy.',.88,6.07,11.2,.46,19,TEAL,True)

    s=slide('Cả lớp thử một câu','CLASSPOINT · DỰ ĐOÁN',True)
    example=data['stages'][2]
    code(s,example['example'],.87,1.75,6.62,3.66,21)
    txt(s,example['question'],7.83,2.05,4.45,1.55,23,WHITE,True)
    txt(s,'Viết đáp án cá nhân. Giải thích với bạn bên cạnh.',.87,5.92,11.2,.5,18,RGBColor(150,225,224))

    s=slide('Đáp án và cách giải thích','CHỐT NHANH')
    card(s,'Đáp án',example['choices'][example['answer']],.88,1.8,4.6,3.2,GREEN,30)
    card(s,'Vì sao?',example['lesson'],5.76,1.8,6.65,3.2,PALE,21)
    txt(s,'Một ví dụ đúng là ví dụ con giải thích được bằng từng bước chạy.',.9,5.83,11.5,.68,21,TEAL,True)

    group=data['group']
    s=slide('Làm mẫu trước khi làm nhóm','BÀI CHUNG · 6 PHÚT')
    txt(s,group['requirement'],.9,1.75,11.45,.95,23,INK,True)
    card(s,'Ví dụ đầu vào',group['sampleIn'],.9,2.98,5.55,2.58,PALE,22)
    card(s,'Ví dụ đầu ra',group['sampleOut'],6.74,2.98,5.55,2.58,GREEN,22)
    txt(s,'Hỏi: máy nhận gì, phải in gì? Chưa cho học sinh code ngay.',.92,6.05,11.6,.5,19,TEAL,True)

    s=slide('Bước 1: Input và Output','LÀM CHUNG VỚI LỚP')
    card(s,'Input',group['input']+'\nVí dụ: '+group['sampleIn'],.9,1.82,5.64,3.9,GREEN,21)
    card(s,'Output',group['output']+'\nVí dụ:\n'+group['sampleOut'],6.79,1.82,5.64,3.9,PALE,21)

    s=slide('Bước 2: mã giả bằng lời','LÀM CHUNG VỚI LỚP')
    card(s,'Mã giả',group['pseudo'],.9,1.71,11.51,4.83,GREEN,21)

    s=slide('Bước 3: đối chiếu với C++','LÀM CHUNG VỚI LỚP')
    code(s,SOLUTION[n],.86,1.63,7.52,5.36,17)
    card(s,'Kiểm tra cùng học sinh','Dòng nào đọc Input?\nDòng nào tạo Output?\nVí dụ còn đúng nếu đổi đầu vào?',8.61,1.7,3.82,4.92,PALE,18)

    s=slide('Đến lượt nhóm','NHÓM 1–3 BẠN · 15 PHÚT')
    card(s,group['name'],group['requirement'],.86,1.71,7.55,3.8,GREEN,21)
    card(s,'Nộp trên Canvas','1. Nhập tên 1–3 bạn.\n2. Ghi Input/Output và mã giả.\n3. Chạy và kiểm tra C++.\n4. Tải một PDF cho nhóm.',8.67,1.71,3.72,3.8,PALE,17)
    txt(s,f'Web nhóm: vuongndlst.github.io/cpp-8/bai{n:02d}/thu-thach.html',.9,5.92,11.5,.62,18,TEAL,True)

    s=slide('Nhóm trình bày và phản hồi','CHIA SẺ · 4 PHÚT',True)
    txt(s,'Một nhóm: nêu Input/Output, mã giả và một dòng C++ quan trọng.',.93,2.05,11.4,1.25,24,WHITE,True)
    txt(s,'Nhóm khác: thử một đầu vào mới và nhận xét kết quả.',.93,3.75,11.4,1.06,24,WHITE,True)
    txt(s,'Giáo viên hỏi: nếu sai, nhóm đã sửa lỗi nào?',.94,5.57,11.3,.61,19,RGBColor(154,227,224))

    s=slide('Tự kiểm tra cuối giờ','CÁ NHÂN · 2 PHÚT')
    card(s,'Việc của con','Nhìn lại mã ở slide làm mẫu.\nChỉ ra dòng quyết định kết quả.\nViết một thay đổi nhỏ và dự đoán đầu ra mới.',.88,1.75,7.4,3.79,GREEN,21)
    card(s,'Thoát học','Nói một ý con nhớ nhất.\nGhi một câu hỏi con còn thắc mắc.',8.55,1.75,3.85,3.79,PALE,19)

    s=slide('Kahoot: chốt bài','3–4 PHÚT',True)
    txt(s,'5 câu ngắn: đọc mã, dự đoán kết quả, chọn cách sửa.',.96,2.09,11.3,1.05,28,WHITE,True)
    txt(s,'Sau câu sai, hỏi học sinh: em đã dựa vào dòng nào để chọn?',.97,3.72,11.1,1.07,23,RGBColor(152,224,220))
    txt(s,'Ghi nhớ: '+data['note'][0],.97,5.42,11.0,.84,20,WHITE)

    path = OUT / f'Slide_Bai{n:02d}_{data["title"].replace(" ", "")}_MachDay_20261003.pptx'
    prs.save(path)
    print(f'Lesson {n}: {len(prs.slides)} slides')
    return path

if __name__ == '__main__':
    for lesson_number in range(2,7):
        make_deck(lesson_number, LESSONS[str(lesson_number)])
